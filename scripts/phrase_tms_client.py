"""Phrase TMS API helpers for glossary sync.

Authenticates with a Phrase Platform API token (OAuth token exchange) and
downloads term-base exports. Used by ``sync_glossaries_from_phrase.py``.
"""

from __future__ import annotations

import json
import os
import re
import urllib.error
import urllib.parse
import urllib.request
from io import BytesIO
from pathlib import Path
from dataclasses import dataclass
from typing import Any

REPO_ROOT = Path(__file__).resolve().parent.parent
TERM_BASES_PATH = Path(__file__).resolve().parent / "phrase_term_bases.json"
SYNC_EXCLUSIONS_PATH = Path(__file__).resolve().parent / "phrase_glossary_sync_exclusions.json"
SYNC_OVERRIDES_PATH = Path(__file__).resolve().parent / "phrase_glossary_sync_overrides.json"

DEFAULT_PLATFORM_BASE_URL = "https://eu.phrase.com"
DEFAULT_TMS_BASE_URL = "https://cloud.memsource.com"


def load_dotenv(path: Path | None = None) -> None:
    """Load ``KEY=value`` pairs from ``.phrase-tms.env`` when present."""
    env_path = path or (REPO_ROOT / ".phrase-tms.env")
    if not env_path.is_file():
        return
    for line in env_path.read_text().splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip())


def platform_token() -> str:
    token = os.environ.get("PHRASE_TMS_TOKEN", "").strip()
    if not token:
        raise RuntimeError(
            "PHRASE_TMS_TOKEN is not set. Copy .phrase-tms.env.example to "
            ".phrase-tms.env or export the variable."
        )
    return token


def platform_base_url() -> str:
    return os.environ.get("PHRASE_PLATFORM_BASE_URL", DEFAULT_PLATFORM_BASE_URL).rstrip("/")


def tms_base_url() -> str:
    return os.environ.get("PHRASE_TMS_BASE_URL", DEFAULT_TMS_BASE_URL).rstrip("/")


def exchange_bearer_jwt(platform_api_token: str | None = None) -> str:
    """Exchange a Platform API token for a short-lived TMS Bearer JWT."""
    token = platform_api_token or platform_token()
    body = urllib.parse.urlencode(
        {
            "grant_type": "urn:ietf:params:oauth:grant-type:token-exchange",
            "subject_token": token,
        }
    ).encode()
    req = urllib.request.Request(
        f"{platform_base_url()}/idm/oauth/token",
        data=body,
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=60) as resp:
        payload = json.load(resp)
    access_token = payload.get("access_token")
    if not access_token:
        raise RuntimeError(f"Token exchange failed: {payload}")
    return access_token


def _api_request(
    bearer: str,
    path: str,
    *,
    method: str = "GET",
    data: bytes | None = None,
    headers: dict[str, str] | None = None,
) -> bytes:
    url = f"{tms_base_url()}/web{path}"
    req_headers = {"Authorization": f"Bearer {bearer}"}
    if headers:
        req_headers.update(headers)
    req = urllib.request.Request(url, data=data, headers=req_headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=300) as resp:
            return resp.read()
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"Phrase TMS {method} {path} failed ({exc.code}): {body}") from exc


def get_term_base(bearer: str, uid: str) -> dict[str, Any]:
    return json.loads(_api_request(bearer, f"/api2/v1/termBases/{uid}"))


def export_term_base_xlsx(bearer: str, uid: str) -> bytes:
    query = urllib.parse.urlencode({"format": "Xlsx", "charset": "UTF-8"})
    return _api_request(
        bearer,
        f"/api2/v1/termBases/{uid}/export?{query}",
        headers={"Accept": "application/octet-stream"},
    )


@dataclass(frozen=True)
class TermBaseConfig:
    """Phrase term-base settings for glossary sync."""

    term_base_uid: str
    term_base_name: str
    source_lang: str
    locales: dict[str, dict[str, str]]


def _normalize_legacy_term_base_config(raw: dict[str, Any]) -> TermBaseConfig:
    """Convert per-locale term base entries to the unified config shape."""
    uids = {entry.get("uid") for entry in raw.values() if isinstance(entry, dict)}
    if len(uids) != 1:
        raise RuntimeError(
            "Legacy phrase_term_bases.json maps locales to different term-base UIDs. "
            "Use the unified term_base + locales structure instead."
        )
    sample = next(entry for entry in raw.values() if isinstance(entry, dict))
    locales = {
        locale_key: {
            "target_lang": entry["target_lang"],
            "source_lang": entry.get("source_lang", "en_us"),
        }
        for locale_key, entry in raw.items()
        if isinstance(entry, dict) and "target_lang" in entry
    }
    source_langs = {
        entry.get("source_lang", "en_us")
        for entry in locales.values()
    }
    if len(source_langs) != 1:
        raise RuntimeError("Legacy phrase_term_bases.json uses mixed source_lang values.")
    return TermBaseConfig(
        term_base_uid=next(iter(uids)),
        term_base_name=sample.get("name", "Phrase term base"),
        source_lang=next(iter(source_langs)),
        locales={
            locale_key: {"target_lang": entry["target_lang"]}
            for locale_key, entry in locales.items()
        },
    )


def load_term_base_config() -> TermBaseConfig:
    raw = json.loads(TERM_BASES_PATH.read_text())
    if "term_base" in raw and "locales" in raw:
        term_base = raw["term_base"]
        return TermBaseConfig(
            term_base_uid=term_base["uid"],
            term_base_name=term_base.get("name", "Phrase term base"),
            source_lang=raw.get("source_lang", "en_us"),
            locales=raw["locales"],
        )
    return _normalize_legacy_term_base_config(raw)


def load_sync_exclusions() -> dict[str, list[str]]:
    """Per-locale English keys to drop after Phrase import (docs-specific guardrails)."""
    if not SYNC_EXCLUSIONS_PATH.is_file():
        return {}
    raw = json.loads(SYNC_EXCLUSIONS_PATH.read_text())
    return {locale: list(keys) for locale, keys in raw.items()}


def normalize_phrase_term(term: str) -> str:
    """Repair Phrase Xlsx export artifacts in term text.

    Some exports insert ``|`` at soft line breaks inside a word (for example
    ``technolog|y``). Remove only pipe characters sandwiched between letters.
    """
    cleaned = re.sub(r"(?<=\w)\|(?=\w)", "", term.strip())
    return re.sub(r"\s+", " ", cleaned)


def load_sync_overrides() -> dict[str, dict[str, str]]:
    """Per-locale English keys to add or replace after Phrase import."""
    if not SYNC_OVERRIDES_PATH.is_file():
        return {}
    raw = json.loads(SYNC_OVERRIDES_PATH.read_text())
    return {
        locale: dict(pairs)
        for locale, pairs in raw.items()
        if isinstance(pairs, dict)
    }


def apply_sync_overrides(
    glossary: dict[str, str],
    locale_key: str,
    overrides: dict[str, dict[str, str]] | None = None,
) -> dict[str, str]:
    """Apply docs-specific glossary corrections on top of Phrase imports."""
    overrides = overrides if overrides is not None else load_sync_overrides()
    locale_overrides = overrides.get(locale_key, {})
    if not locale_overrides:
        return glossary
    merged = dict(glossary)
    merged.update(locale_overrides)
    return merged


def apply_sync_exclusions(
    glossary: dict[str, str],
    locale_key: str,
    exclusions: dict[str, list[str]] | None = None,
) -> dict[str, str]:
    """Remove locale-specific keys that are unsafe for docs glossary matching."""
    exclusions = exclusions if exclusions is not None else load_sync_exclusions()
    skip = set(exclusions.get(locale_key, []))
    if not skip:
        return glossary
    return {key: value for key, value in glossary.items() if key not in skip}


def parse_term_base_xlsx(
    xlsx_bytes: bytes,
    *,
    source_lang: str,
    target_lang: str,
) -> dict[str, str]:
    """Convert a Phrase term-base Xlsx export to ``{english: translation}``."""
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required. Install with: pip install -r scripts/requirements-glossaries.txt"
        ) from exc

    workbook = openpyxl.load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return {}

    source_cols = [i for i, label in enumerate(header) if label == source_lang]
    target_cols = [i for i, label in enumerate(header) if label == target_lang]
    if not source_cols or not target_cols:
        raise RuntimeError(
            f"Xlsx header missing {source_lang!r} or {target_lang!r} columns: {header}"
        )

    glossary: dict[str, str] = {}
    for row in rows:
        if not row:
            continue
        source_terms = []
        for col in source_cols:
            if col < len(row) and row[col]:
                term = normalize_phrase_term(str(row[col]))
                if term and term not in source_terms:
                    source_terms.append(term)
        target_terms = []
        for col in target_cols:
            if col < len(row) and row[col]:
                term = normalize_phrase_term(str(row[col]))
                if term and term not in target_terms:
                    target_terms.append(term)
        if not source_terms or not target_terms:
            continue
        translation = " or ".join(target_terms)
        for english in source_terms:
            glossary[english] = translation
    workbook.close()
    return glossary


def parse_term_base_xlsx_multi(
    xlsx_bytes: bytes,
    *,
    source_lang: str,
    locale_targets: dict[str, str],
) -> dict[str, dict[str, str]]:
    """Parse one Phrase Xlsx export into per-locale ``{english: translation}`` maps."""
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required. Install with: pip install -r scripts/requirements-glossaries.txt"
        ) from exc

    workbook = openpyxl.load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        workbook.close()
        return {locale_key: {} for locale_key in locale_targets}

    source_cols = [i for i, label in enumerate(header) if label == source_lang]
    target_cols = {
        locale_key: [i for i, label in enumerate(header) if label == target_lang]
        for locale_key, target_lang in locale_targets.items()
    }
    missing_targets = [
        locale_key
        for locale_key, cols in target_cols.items()
        if not cols
    ]
    if not source_cols:
        workbook.close()
        raise RuntimeError(
            f"Xlsx header missing source column {source_lang!r}: {header}"
        )
    if missing_targets:
        workbook.close()
        raise RuntimeError(
            "Xlsx header missing target column(s) for locale(s) "
            f"{missing_targets}: {header}"
        )

    glossaries = {locale_key: {} for locale_key in locale_targets}
    for row in rows:
        if not row:
            continue
        source_terms = []
        for col in source_cols:
            if col < len(row) and row[col]:
                term = normalize_phrase_term(str(row[col]))
                if term and term not in source_terms:
                    source_terms.append(term)
        if not source_terms:
            continue
        for locale_key, cols in target_cols.items():
            target_terms = []
            for col in cols:
                if col < len(row) and row[col]:
                    term = normalize_phrase_term(str(row[col]))
                    if term and term not in target_terms:
                        target_terms.append(term)
            if not target_terms:
                continue
            translation = " or ".join(target_terms)
            for english in source_terms:
                glossaries[locale_key][english] = translation
    workbook.close()
    return glossaries


def fetch_glossaries_for_locales(
    bearer: str,
    locale_keys: list[str],
    config: TermBaseConfig | None = None,
) -> dict[str, dict[str, str]]:
    """Download the configured term base once and parse selected locales."""
    config = config or load_term_base_config()
    unknown = [locale_key for locale_key in locale_keys if locale_key not in config.locales]
    if unknown:
        raise KeyError(
            f"Unknown locale(s): {unknown!r}; expected one of {sorted(config.locales)}"
        )
    xlsx_bytes = export_term_base_xlsx(bearer, config.term_base_uid)
    locale_targets = {
        locale_key: config.locales[locale_key]["target_lang"]
        for locale_key in locale_keys
    }
    glossaries = parse_term_base_xlsx_multi(
        xlsx_bytes,
        source_lang=config.source_lang,
        locale_targets=locale_targets,
    )
    return {
        locale_key: apply_sync_overrides(
            apply_sync_exclusions(glossaries[locale_key], locale_key),
            locale_key,
        )
        for locale_key in locale_keys
    }


def fetch_glossary_for_locale(
    bearer: str,
    locale_key: str,
    config: TermBaseConfig | None = None,
) -> dict[str, str]:
    """Download and parse one locale from the configured term base."""
    return fetch_glossaries_for_locales(bearer, [locale_key], config)[locale_key]
